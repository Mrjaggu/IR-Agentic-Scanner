"""
brief_export.py -- Section: Prep brief, exported as a real Word doc or PDF,
not just markdown.

Motivation: brief_as_markdown() (src/data/workspace_state.py) has always
been the export path, but a raw .md file is a poor deliverable for anyone
who isn't going to open it in an editor -- an IR lead sharing the call-day
brief with the team, or printing it, wants a Word doc or a PDF, not
markdown syntax. Both renderers below build from the SAME grouped-items
data workspace_state.grouped_brief_items() already produces for markdown,
so all three formats stay structurally in sync -- add a new brief-item
kind once, get it rendered correctly in all three, not just the one
someone remembered to update.

Both libraries (python-docx, reportlab) are already dependencies elsewhere
in the Python ecosystem this platform runs on and needed no new install --
kept optional here anyway (import inside each function, not at module
level) so a future environment missing one doesn't break the whole module
just because someone imported it for the other format.
"""

from datetime import datetime, timezone


def brief_as_docx(quarter_label: str = "", bank_id: str = "axis") -> bytes:
    """Returns the .docx file's raw bytes -- the route wraps this in a
    StreamingResponse, nothing here touches disk."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from src.data.workspace_state import grouped_brief_items, KIND_LABEL

    doc = Document()
    title = f"Prep brief{' — ' + quarter_label if quarter_label else ''}"
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}").italic = True

    grouped = grouped_brief_items(bank_id)
    if not grouped:
        doc.add_paragraph("Nothing pinned yet.")
        return _docx_bytes(doc)

    for kind, items in grouped:
        doc.add_heading(KIND_LABEL.get(kind, kind), level=1)
        for it in items:
            p = doc.add_paragraph()
            run = p.add_run(it["title"])
            run.bold = True
            doc.add_paragraph(it["body"])
            meta = it.get("meta") or {}
            tag_bits = [f"{k}: {v}" for k, v in meta.items() if v]
            if tag_bits:
                tag_p = doc.add_paragraph(" · ".join(tag_bits))
                for r in tag_p.runs:
                    r.italic = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return _docx_bytes(doc)


def _docx_bytes(doc) -> bytes:
    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def brief_as_pdf(quarter_label: str = "", bank_id: str = "axis") -> bytes:
    """Returns the .pdf file's raw bytes. reportlab's Platypus flowables
    handle pagination automatically -- a long brief just spans more pages,
    nothing here has to compute page breaks by hand."""
    import io
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from xml.sax.saxutils import escape as _xml_escape
    from src.data.workspace_state import grouped_brief_items, KIND_LABEL

    buf = io.BytesIO()
    docpdf = SimpleDocTemplate(buf, pagesize=LETTER,
                               leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                               topMargin=0.9 * inch, bottomMargin=0.9 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("BriefTitle", parent=styles["Title"], fontSize=20, spaceAfter=4)
    meta_style = ParagraphStyle("BriefMeta", parent=styles["Normal"], fontSize=9,
                                textColor=HexColor("#666666"), spaceAfter=16)
    h1_style = ParagraphStyle("BriefH1", parent=styles["Heading1"], fontSize=14, spaceBefore=14, spaceAfter=8)
    item_title_style = ParagraphStyle("BriefItemTitle", parent=styles["Normal"], fontSize=11,
                                      leading=14, spaceAfter=2, fontName="Helvetica-Bold")
    body_style = ParagraphStyle("BriefBody", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=4)
    tag_style = ParagraphStyle("BriefTag", parent=styles["Normal"], fontSize=8.5,
                               textColor=HexColor("#888888"), spaceAfter=12)

    def esc(s):
        return _xml_escape(str(s or "")).replace("\n", "<br/>")

    flow = []
    title = f"Prep brief{' — ' + quarter_label if quarter_label else ''}"
    flow.append(Paragraph(esc(title), title_style))
    flow.append(Paragraph(esc(f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"), meta_style))

    grouped = grouped_brief_items(bank_id)
    if not grouped:
        flow.append(Paragraph("Nothing pinned yet.", body_style))
    else:
        for kind, items in grouped:
            flow.append(Paragraph(esc(KIND_LABEL.get(kind, kind)), h1_style))
            for it in items:
                flow.append(Paragraph(esc(it["title"]), item_title_style))
                flow.append(Paragraph(esc(it["body"]), body_style))
                meta = it.get("meta") or {}
                tag_bits = [f"{k}: {v}" for k, v in meta.items() if v]
                if tag_bits:
                    flow.append(Paragraph(esc(" · ".join(tag_bits)), tag_style))
                else:
                    flow.append(Spacer(1, 8))

    docpdf.build(flow)
    return buf.getvalue()
